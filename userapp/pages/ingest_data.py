import streamlit as st
import pandas as pd
import json
import sys
import os
import datetime
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from schema import ENTITY_SCHEMAS, VOCABULARIES
from api_client import (
    create_project, create_sample, create_experiment,
    create_run, create_flowcell_library, create_file,
    create_cohort, create_cohort_member,
    create_pool, create_pool_member, create_subject, create_sample_source
)
import math

# ── Ingest audit trail ──────────────────────────────────────────────────────
# Archives every successfully-submitted ingest file (or a snapshot of manual
# entries) to GCS for posterity. Uses the same secret-configured bucket
# pattern as the group template storage below. Failure here never blocks
# an ingest — it's a best-effort archive, not a transactional requirement.

def _get_audit_bucket():
    try:
        return st.secrets.get("AUDIT_BUCKET") or os.getenv("AUDIT_BUCKET")
    except Exception:
        return os.getenv("AUDIT_BUCKET")


def audit_archive(entity: str, file_bytes: bytes, filename: str, row_count: int, source: str):
    """Best-effort upload of an ingested file to the audit bucket.
    Returns True if archived, False if skipped or failed (never raises)."""
    bucket_name = _get_audit_bucket()
    if not bucket_name:
        return False

    try:
        from google.cloud import storage as gcs_lib

        client = gcs_lib.Client()
        bucket = client.bucket(bucket_name)

        now = datetime.datetime.utcnow()
        date_str = now.strftime("%Y-%m-%d")
        time_str = now.strftime("%H-%M-%S")
        blob_path = f"ingest_audit/{entity}/{date_str}/{time_str}_{filename}"

        blob = bucket.blob(blob_path)
        blob.metadata = {
            "entity":       entity,
            "source":       source,  # "csv_upload" | "xlsx_upload" | "manual_entry"
            "row_count":    str(row_count),
            "ingested_at":  now.isoformat(),
        }
        blob.upload_from_string(file_bytes)
        return True
    except Exception as e:
        st.warning(f"Audit archive skipped: {e}")
        return False

API_FN = {
    "projects":          create_project,
    "subjects":          create_subject,
    "samples":           create_sample,
    "sample_sources":    create_sample_source,
    "cohorts":           create_cohort,
    "cohort_members":    create_cohort_member,
    "experiments":       create_experiment,
    "pools":             create_pool,
    "pool_members":      create_pool_member,
    "sequencing_runs":   create_run,
    "flowcell_libraries": create_flowcell_library,
    "files":             create_file,
}

# ── Template system ───────────────────────────────────────────────────────────

EXAMPLE_VALUES = {
    "project_name":      "IBD_Atlas_2026",
    "description":       "Inflammatory bowel disease single-cell atlas",
    "sample_name":       "SAMPLE-001",
    "organism":          "Homo sapiens",
    "tissue":            "colon",
    "sample_type":       "individual",
    "pub_id":            "DONOR-001",
    "freezerworks_id":   "FW-12345",
    "assay_type":        "scRNA-seq",
    "library_prep_date": "2026-06-01",
    "library_protocol":  "10x Chromium v3",
    "library_version":   "GEX v3.1",
    "flowcell_id":       "HXXXXDRXY",
    "machine":           "NovaSeq 6000",
    "run_date":          "2026-06-15",
    "read_length":       "28x90",
    "sequencing_center": "Broad Genomics Platform",
    "bcl_gcs_uri":       "gs://my-bucket/run/HXXXXDRXY/",
    "lane":              "1",
    "index_sequence":    "ATCGATCG",
    "cohort_name":       "Healthy_Controls",
    "cohort_type":       "biological",
    "pool_name":         "Pool_A",
    "gcs_uri":           "gs://my-bucket/samples/SAMPLE-001.fastq.gz",
    "gcs_bucket":        "my-bucket",
    "file_path":         "/path/to/SAMPLE-001.fastq.gz",
    "file_type":         "fastq",
    "file_format":       "fastq.gz",
    "subject_id":        "XSU00001",
}

TEMPLATE_CATALOG = [
    # ── Single-entity ──────────────────────────────────────────────────────
    {"name": "Projects",           "file": "xavier_template_projects.csv",
     "desc": "Register new top-level projects or studies.",
     "entities": ["projects"], "workflow": False},
    {"name": "Subjects",           "file": "xavier_template_subjects.csv",
     "desc": "Register donors, patients, or animal subjects.",
     "entities": ["subjects"], "workflow": False},
    {"name": "Samples",            "file": "xavier_template_samples.csv",
     "desc": "Register biospecimens under an existing project.",
     "entities": ["samples"], "workflow": False},
    {"name": "Sample Sources",     "file": "xavier_template_sample_sources.csv",
     "desc": "Link subjects to samples (one row per subject–sample pair).",
     "entities": ["sample_sources"], "workflow": False},
    {"name": "Experiments",        "file": "xavier_template_experiments.csv",
     "desc": "Register library preparations under an existing sample.",
     "entities": ["experiments"], "workflow": False},
    {"name": "Sequencing Runs",    "file": "xavier_template_sequencing_runs.csv",
     "desc": "Register sequencing flowcells.",
     "entities": ["sequencing_runs"], "workflow": False},
    {"name": "Flowcell Libraries", "file": "xavier_template_flowcell_libraries.csv",
     "desc": "Register per-lane library assignments for a sequencing run.",
     "entities": ["flowcell_libraries"], "workflow": False},
    {"name": "Files",              "file": "xavier_template_files.csv",
     "desc": "Register output files (FASTQ, BAM, count matrices).",
     "entities": ["files"], "workflow": False},
    {"name": "Cohorts",            "file": "xavier_template_cohorts.csv",
     "desc": "Register sample groupings (biological, technical, analysis).",
     "entities": ["cohorts"], "workflow": False},
    {"name": "Cohort Members",     "file": "xavier_template_cohort_members.csv",
     "desc": "Link samples to an existing cohort.",
     "entities": ["cohort_members"], "workflow": False},
    {"name": "Pools",              "file": "xavier_template_pools.csv",
     "desc": "Register library pool groups.",
     "entities": ["pools"], "workflow": False},
    {"name": "Pool Members",       "file": "xavier_template_pool_members.csv",
     "desc": "Link experiments to an existing pool.",
     "entities": ["pool_members"], "workflow": False},
    # ── Multi-step workflows ───────────────────────────────────────────────
    {"name": "Projects + Samples",
     "file": "xavier_template_wf_projects_samples.csv",
     "desc": "Onboard a new project with its first samples.",
     "entities": ["projects", "samples"], "workflow": True,
     "steps": ["1. Upload as Projects", "2. Upload as Samples"]},
    {"name": "Subjects + Samples + Sample Sources",
     "file": "xavier_template_wf_subjects_samples_sources.csv",
     "desc": "Register subjects, their samples, and subject–sample links together.",
     "entities": ["subjects", "samples", "sample_sources"], "workflow": True,
     "steps": ["1. Upload as Subjects", "2. Upload as Samples", "3. Upload as Sample Sources"]},
    {"name": "Samples + Experiments",
     "file": "xavier_template_wf_samples_experiments.csv",
     "desc": "Register samples and their library preparations in one sheet.",
     "entities": ["samples", "experiments"], "workflow": True,
     "steps": ["1. Upload as Samples", "2. Upload as Experiments"]},
    {"name": "Sequencing Run + Flowcell Libraries",
     "file": "xavier_template_wf_run_libraries.csv",
     "desc": "Register a flowcell with all per-lane library assignments.",
     "entities": ["sequencing_runs", "flowcell_libraries"], "workflow": True,
     "steps": ["1. Upload as Sequencing Runs", "2. Upload as Flowcell Libraries"]},
    {"name": "Full Prep Manifest",
     "file": "xavier_template_wf_full_prep.csv",
     "desc": "Projects → Samples → Experiments in one planning sheet.",
     "entities": ["projects", "samples", "experiments"], "workflow": True,
     "steps": ["1. Upload as Projects", "2. Upload as Samples", "3. Upload as Experiments"]},
]


def _template_columns(entities):
    """Return deduplicated ordered column list for the given entities."""
    cols, seen = [], set()
    for ent in entities:
        schema = ENTITY_SCHEMAS[ent]
        required = schema["required"]
        optional = schema["optional"]
        for col in required + optional:
            if col not in seen:
                cols.append((col, col in required))
                seen.add(col)
    return cols  # list of (col_name, is_required)


def generate_template_csv(entities):
    """CSV template with one example row."""
    cols = [c for c, _ in _template_columns(entities)]
    example = {col: EXAMPLE_VALUES.get(col, "") for col in cols}
    return pd.DataFrame([example]).to_csv(index=False).encode("utf-8")


def generate_template_xlsx(entities):
    """XLSX template with styled headers, one example row, and dropdown validation."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.utils import get_column_letter
        from io import BytesIO
    except ImportError:
        return None  # caller falls back to CSV

    wb = Workbook()
    ws = wb.active
    ws.title = "Template"

    cols = _template_columns(entities)

    FILL_REQUIRED = PatternFill("solid", fgColor="D6E4F7")  # light blue
    FILL_OPTIONAL = PatternFill("solid", fgColor="F0F0F0")  # light grey
    FILL_EXAMPLE  = PatternFill("solid", fgColor="EAF7EA")  # light green

    for i, (col, is_required) in enumerate(cols, start=1):
        letter = get_column_letter(i)

        # Header cell
        hdr = ws.cell(row=1, column=i, value=col)
        hdr.font      = Font(bold=True, size=11)
        hdr.fill      = FILL_REQUIRED if is_required else FILL_OPTIONAL
        hdr.alignment = Alignment(horizontal="center")

        # Example cell
        ex = ws.cell(row=2, column=i, value=EXAMPLE_VALUES.get(col, ""))
        ex.fill      = FILL_EXAMPLE
        ex.font      = Font(italic=True, color="666666")

        # Column width
        ws.column_dimensions[letter].width = max(len(col) + 4, 16)

        # Dropdown validation for vocabulary fields
        if col in VOCABULARIES:
            options = ",".join(VOCABULARIES[col])
            dv = DataValidation(
                type="list",
                formula1=f'"{options}"',
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Invalid value",
                error=f"Choose from: {options}",
            )
            ws.add_data_validation(dv)
            dv.sqref = f"{letter}3:{letter}5000"

    # Freeze header row
    ws.freeze_panes = "A3"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


st.title("📥 Ingest Data")

tab1, tab2, tab3 = st.tabs(["Upload CSV", "Manual Entry", "Templates"])

# ── Helper functions ──────────────────────────────────────────────────────────

def sanitize_payload(obj):
    if isinstance(obj, dict):
        return {k: sanitize_payload(v) for k, v in obj.items()}
    elif isinstance(obj, list):
        return [sanitize_payload(v) for v in obj]
    elif isinstance(obj, float):
        if math.isnan(obj) or math.isinf(obj):
            return None
    return obj


def build_payload(row_dict, required_columns, optional_columns, no_metadata=False):
    core = {}
    metadata = {}
    all_core = required_columns + optional_columns

    for k, v in row_dict.items():
        v = None if pd.isna(v) else v

        if k in all_core:
            core[k] = v
        elif not no_metadata and v not in ["", None]:
            metadata[k] = v

    if not no_metadata:
        core["extra_metadata"] = metadata

    return sanitize_payload(core)


def validate_payload(payload, required_columns):
    errors = []
    for col in required_columns:
        if col not in payload or payload[col] in [None, ""]:
            errors.append(f"Missing required field: {col}")
    return errors


def schema_helper(entity):
    """Renders required / optional field info for the selected entity."""
    schema = ENTITY_SCHEMAS[entity]
    required = schema["required"]
    optional = schema["optional"]

    col1, col2 = st.columns(2)

    with col1:
        st.markdown("**Required fields**")
        for f in required:
            st.markdown(f"- `{f}`")

    with col2:
        st.markdown("**Optional fields**")
        if optional:
            for f in optional:
                st.markdown(f"- `{f}`")
        else:
            st.markdown("_None_")

    if not schema.get("no_metadata"):
        st.caption("Any additional columns in your CSV will be stored in extra_metadata.")

# ── CSV Upload ────────────────────────────────────────────────────────────────

with tab1:
    st.subheader("Upload CSV")

    entity = st.selectbox(
        "Entity type",
        list(ENTITY_SCHEMAS.keys()),
        key="entity_tab1"
    )

    schema        = ENTITY_SCHEMAS[entity]
    required_cols = schema["required"]
    optional_cols = schema["optional"]
    no_metadata   = schema.get("no_metadata", False)
    api_fn        = API_FN[entity]

    schema_helper(entity)
    st.divider()

    uploaded_file = st.file_uploader("Upload CSV or XLSX", type=["csv", "xlsx"])

    if uploaded_file:
        is_xlsx    = uploaded_file.name.endswith(".xlsx")
        file_bytes = uploaded_file.getvalue()  # capture before pandas reads the stream

        if is_xlsx:
            df = pd.read_excel(uploaded_file, dtype=str).fillna("")
        else:
            df = pd.read_csv(uploaded_file, dtype=str).fillna("")

        st.write("Preview:")
        st.dataframe(df)

        if st.button("Validate CSV"):
            all_errors = []

            for i, row in df.iterrows():
                payload = build_payload(
                    row.to_dict(), required_cols, optional_cols, no_metadata
                )
                errors = validate_payload(payload, required_cols)

                if errors:
                    all_errors.append({"row": i, "errors": errors})

            if all_errors:
                st.error("Validation errors found")
                st.json(all_errors)
            else:
                st.success("CSV validation passed")

        if st.button("Submit all rows"):
            success = 0
            errors  = []

            for i, row in df.iterrows():
                payload = build_payload(
                    row.to_dict(), required_cols, optional_cols, no_metadata
                )

                try:
                    api_fn(payload)
                    success += 1
                except Exception as e:
                    errors.append({"row": i, "error": str(e)})

            st.success(f"{success} rows inserted")

            if errors:
                st.error("Errors:")
                st.json(errors)

            if success > 0:
                archived = audit_archive(
                    entity, file_bytes, uploaded_file.name, success,
                    source="xlsx_upload" if is_xlsx else "csv_upload"
                )
                if archived:
                    st.caption("📦 Original file archived for audit.")


# ── Manual Entry ──────────────────────────────────────────────────────────────

with tab2:

    entity = st.selectbox(
        "Entity type",
        list(ENTITY_SCHEMAS.keys()),
        key="entity_tab2"
    )

    schema        = ENTITY_SCHEMAS[entity]
    required_cols = schema["required"]
    optional_cols = schema["optional"]
    no_metadata   = schema.get("no_metadata", False)
    api_fn        = API_FN[entity]

    schema_helper(entity)
    st.divider()

    if "metadata_fields" not in st.session_state:
        st.session_state.metadata_fields = []

    if not no_metadata:
        st.subheader("Extra metadata fields")
        new_field = st.text_input("Add metadata field")

        if st.button("Add field"):
            if new_field and new_field not in st.session_state.metadata_fields:
                st.session_state.metadata_fields.append(new_field)

    all_columns  = required_cols + optional_cols
    if not no_metadata:
        all_columns += st.session_state.metadata_fields

    df_template = pd.DataFrame([{col: "" for col in all_columns}])

    edited_df = st.data_editor(
        df_template,
        num_rows="dynamic",
        use_container_width=True,
        column_config={
            col: st.column_config.TextColumn(
                col,
                help="Required" if col in required_cols else "Optional"
            )
            for col in all_columns
        }
    )

    if st.button("Validate"):
        errors = []

        for i, row in edited_df.iterrows():
            row_dict = row.to_dict()

            if all(v in ["", None] for v in row_dict.values()):
                continue

            payload    = build_payload(row_dict, required_cols, optional_cols, no_metadata)
            row_errors = validate_payload(payload, required_cols)

            if row_errors:
                errors.append({"row": i, "errors": row_errors})

        if errors:
            st.error("Validation issues found")
            st.json(errors)
        else:
            st.success("All rows valid")

    if st.button("Submit data"):
        success = 0
        errors  = []

        for i, row in edited_df.iterrows():
            row_dict = row.to_dict()

            if all(v in ["", None] for v in row_dict.values()):
                continue

            payload = build_payload(row_dict, required_cols, optional_cols, no_metadata)

            try:
                api_fn(payload)
                success += 1
            except Exception as e:
                errors.append({"row": i, "error": str(e)})

        st.success(f"{success} rows inserted")

        if errors:
            st.error("Some rows failed")
            st.json(errors)

        if success > 0:
            snapshot_bytes = edited_df.to_csv(index=False).encode("utf-8")
            archived = audit_archive(
                entity, snapshot_bytes, f"manual_entry_{entity}.csv", success,
                source="manual_entry"
            )
            if archived:
                st.caption("📦 Submission snapshot archived for audit.")


# ── Templates ─────────────────────────────────────────────────────────────────

with tab3:
    st.subheader("Download Templates")
    st.caption(
        "Each template is generated from the current schema and includes one example row. "
        "Download, fill in your data, then upload in the **Upload CSV** tab. "
        "Multi-step templates combine columns from several entity types — "
        "upload each section separately in the order shown."
    )

    # ── Single-entity templates ──────────────────────────────────────────────
    st.markdown("### Single-entity templates")

    single = [t for t in TEMPLATE_CATALOG if not t["workflow"]]
    for i in range(0, len(single), 3):
        row_tmpl = single[i : i + 3]
        cols = st.columns(3)
        for j, tmpl in enumerate(row_tmpl):
            with cols[j]:
                st.markdown(f"**{tmpl['name']}**")
                st.caption(tmpl["desc"])
                xlsx_bytes = generate_template_xlsx(tmpl["entities"])
                bc1, bc2 = st.columns(2)
                with bc1:
                    st.download_button(
                        label="⬇ XLSX",
                        data=xlsx_bytes or generate_template_csv(tmpl["entities"]),
                        file_name=tmpl["file"].replace(".csv", ".xlsx") if xlsx_bytes else tmpl["file"],
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if xlsx_bytes else "text/csv",
                        key=f"dlx_{tmpl['file']}",
                        use_container_width=True,
                    )
                with bc2:
                    st.download_button(
                        label="⬇ CSV",
                        data=generate_template_csv(tmpl["entities"]),
                        file_name=tmpl["file"],
                        mime="text/csv",
                        key=f"dlc_{tmpl['file']}",
                        use_container_width=True,
                    )

    st.divider()

    # ── Workflow templates ───────────────────────────────────────────────────
    st.markdown("### Multi-step workflow templates")
    st.caption(
        "Fill the full sheet end-to-end, then split and upload each section "
        "as its own entity type in the order listed."
    )

    for tmpl in [t for t in TEMPLATE_CATALOG if t["workflow"]]:
        with st.expander(f"**{tmpl['name']}**  —  {tmpl['desc']}"):
            c1, c2 = st.columns([2, 1])
            with c1:
                st.markdown("**Entities covered (in upload order):**")
                for step in tmpl.get("steps", []):
                    st.markdown(f"- {step}")
                entity_labels = " → ".join(
                    ENTITY_SCHEMAS[e]["label"] for e in tmpl["entities"]
                )
                st.caption(f"Schema: {entity_labels}")
            with c2:
                xlsx_bytes = generate_template_xlsx(tmpl["entities"])
                st.download_button(
                    label="⬇ XLSX",
                    data=xlsx_bytes or generate_template_csv(tmpl["entities"]),
                    file_name=tmpl["file"].replace(".csv", ".xlsx") if xlsx_bytes else tmpl["file"],
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" if xlsx_bytes else "text/csv",
                    key=f"dlx_{tmpl['file']}",
                    use_container_width=True,
                )
                st.download_button(
                    label="⬇ CSV",
                    data=generate_template_csv(tmpl["entities"]),
                    file_name=tmpl["file"],
                    mime="text/csv",
                    key=f"dlc_{tmpl['file']}",
                    use_container_width=True,
                )

    # ── GCS custom templates (optional) ─────────────────────────────────────
    try:
        gcs_bucket = st.secrets.get("TEMPLATE_BUCKET") or os.getenv("TEMPLATE_BUCKET")
    except Exception:
        gcs_bucket = os.getenv("TEMPLATE_BUCKET")

    if gcs_bucket:
        st.divider()
        st.markdown("### Group templates")
        st.caption(
            "Templates uploaded by lab groups, organized by group name. "
            "These can include extra_metadata columns specific to your group's workflow."
        )

        try:
            from google.cloud import storage as gcs_lib
            import datetime

            client = gcs_lib.Client()
            bucket = client.bucket(gcs_bucket)

            # ── Upload form ──────────────────────────────────────────────────
            with st.expander("Upload a template", icon="📤"):
                with st.form("upload_template"):
                    group_name   = st.text_input(
                        "Group name",
                        placeholder="e.g. GI, Dermatology, Shared"
                    )
                    tmpl_desc    = st.text_input(
                        "Description",
                        placeholder="What is this template for?"
                    )
                    upload_file  = st.file_uploader(
                        "Template file", type=["csv", "xlsx"]
                    )
                    do_upload    = st.form_submit_button("Upload")

                if do_upload:
                    if not group_name.strip():
                        st.error("Group name is required.")
                    elif upload_file is None:
                        st.error("Select a file to upload.")
                    else:
                        safe_group = group_name.strip().replace(" ", "_")
                        blob_path  = f"templates/{safe_group}/{upload_file.name}"
                        blob       = bucket.blob(blob_path)
                        blob.metadata = {
                            "description":  tmpl_desc.strip(),
                            "uploaded_at":  datetime.datetime.utcnow().isoformat(),
                        }
                        blob.upload_from_file(upload_file, rewind=True)
                        st.success(
                            f"Uploaded **{upload_file.name}** to group **{safe_group}**."
                        )
                        st.rerun()

            # ── Browse by group ──────────────────────────────────────────────
            all_blobs = [
                b for b in bucket.list_blobs(prefix="templates/")
                if b.name.endswith((".csv", ".xlsx"))
            ]

            if not all_blobs:
                st.info("No group templates uploaded yet.")
            else:
                # Parse group from path: templates/{group}/{filename}
                groups = {}
                for blob in all_blobs:
                    parts = blob.name.split("/")
                    if len(parts) >= 3:
                        grp = parts[1].replace("_", " ")
                        groups.setdefault(grp, []).append(blob)

                for grp, blobs in sorted(groups.items()):
                    st.markdown(f"**{grp}**")
                    for blob in blobs:
                        fname    = blob.name.split("/")[-1]
                        meta     = blob.metadata or {}
                        desc     = meta.get("description", "")
                        uploaded = meta.get("uploaded_at", "")[:10] if meta.get("uploaded_at") else ""
                        mime     = (
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                            if fname.endswith(".xlsx") else "text/csv"
                        )
                        c1, c2 = st.columns([3, 1])
                        with c1:
                            st.caption(f"`{fname}`" + (f"  —  {desc}" if desc else "") +
                                       (f"  ·  {uploaded}" if uploaded else ""))
                        with c2:
                            st.download_button(
                                label="⬇ Download",
                                data=blob.download_as_bytes(),
                                file_name=fname,
                                mime=mime,
                                key=f"gcs_{blob.name}",
                                use_container_width=True,
                            )
                    st.divider()

        except ImportError:
            st.warning(
                "google-cloud-storage is not installed — group templates unavailable."
            )
        except Exception as e:
            st.warning(f"Could not connect to template storage: {e}")